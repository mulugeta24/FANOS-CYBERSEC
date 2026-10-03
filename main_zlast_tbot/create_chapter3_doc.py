#!/usr/bin/env python3
"""
Generate Updated Chapter 3 Documentation as Word Document
Includes all new features and improvements from the project
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

def main():
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Title Page
    title = doc.add_heading('Chapter 3', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_heading('Malicious Link Detector Bot', level=1)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Date
    date_para = doc.add_paragraph(f'Updated: {datetime.now().strftime("%B %d, %Y")}')
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    # Section 3.1: Summary
    doc.add_heading('3.1. Summary of the Project', level=2)
    
    doc.add_paragraph(
        'This project presents the design, development, and implementation of a comprehensive hybrid '
        'malicious URL detection system. The primary objective was to create a robust, scalable, and '
        'accessible tool capable of automatically identifying and classifying URLs that pose a threat to '
        'users, such as phishing, malware, and defacement sites. The solution is integrated with multiple '
        'platforms: a Telegram bot for instant messaging-based analysis, a web-based admin dashboard for '
        'monitoring and management, and a unified backend API serving both interfaces.'
    )
    
    doc.add_paragraph(
        'The core of the system is a two-tiered, hybrid detection model. The first tier leverages a machine '
        'learning classifier, specifically a trained XGBoost model, to analyze a URL\'s intrinsic features, '
        'such as its length, number of special characters, subdomain count, and presence of suspicious keywords. '
        'This model provides a fast, initial assessment of the URL\'s risk profile with 17 carefully engineered '
        'features. The second tier integrates with the VirusTotal API, a reputable external service that provides '
        'real-time, comprehensive threat intelligence by scanning URLs against a vast repository of security '
        'vendors and blocklists.'
    )
    
    doc.add_paragraph(
        'The system\'s intelligent verdict synthesis logic prioritizes external threat intelligence from '
        'VirusTotal while using the machine learning model as a reliable fallback or supplementary tool. A key '
        'innovation is the confidence-aware verdict system: when ML confidence falls below the 85% threshold and '
        'VirusTotal confirms a URL is safe, the system trusts VirusTotal over uncertain ML predictions, '
        'significantly reducing false positives. To enhance performance and reduce redundant API calls, a SQLite '
        'database was implemented with separate tables for automatic detections (scans) and user feedback '
        '(user_reports), providing a scalable and efficient solution with proper data separation for production use.'
    )
    
    doc.add_paragraph(
        'The project successfully delivers a functional, modular, and production-ready system that effectively '
        'addresses the problem of malicious URLs. Its hybrid approach ensures high accuracy, the Telegram interface '
        'makes it a practical security tool for everyday users, and the web dashboard provides administrators with '
        'comprehensive monitoring and management capabilities.'
    )
    
    doc.add_page_break()
    
    # Section 3.2: Problem Statement
    doc.add_heading('3.2. Problem Statement & Justification', level=2)
    
    doc.add_heading('3.2.1. Problem Statement', level=3)
    
    doc.add_paragraph(
        'In the modern digital landscape, the proliferation of malicious URLs—including phishing links, malware '
        'distribution sites, and defaced web pages—represents a significant and growing threat to individuals and '
        'organizations. These threats exploit human trust and vulnerabilities to steal sensitive information, '
        'distribute harmful software, and compromise online security. The sheer volume and dynamic nature of these '
        'malicious URLs make it nearly impossible for a user to manually verify their safety. The existing system, '
        'often reliant on human-driven reporting and slow, reactive blocklist updates, is insufficient to provide '
        'timely protection against new and evolving threats. This reactive approach leaves a critical window of '
        'opportunity for attackers to cause significant harm before a URL is officially flagged as malicious. '
        'Therefore, there is a clear need for a proactive, automated, and accessible solution that can provide '
        'real-time threat analysis of URLs to prevent user exposure to these risks.'
    )
    
    doc.add_heading('3.2.2. Justification (Based on the Current/Existing System)', level=3)
    
    doc.add_paragraph(
        'The current/existing system for URL threat detection primarily operates on a reactive model. When a '
        'malicious URL is discovered, it is manually reported to security vendors. These vendors then analyze the '
        'URL, add it to their blacklists, and update browser-based security features. This process, however, has '
        'several limitations:'
    )
    
    doc.add_paragraph('Latency: The time lag between a new malicious URL\'s creation and its inclusion in public '
                     'blocklists can range from hours to days. During this period, the URL can actively be used to '
                     'compromise unsuspecting users.', style='List Bullet')
    
    doc.add_paragraph('Lack of Accessibility: Some security tools are integrated into specific browsers or require '
                     'technical expertise to use, limiting accessibility for the general population.', 
                     style='List Bullet')
    
    doc.add_paragraph('Human Error: Users may overlook subtle signs of a malicious URL (e.g., misspellings, unusual '
                     'characters) or may not have the knowledge to identify them.', style='List Bullet')
    
    doc.add_paragraph('Inefficiency: The manual reporting and updating system is inefficient for handling the vast '
                     'number of new malicious URLs that appear daily.', style='List Bullet')
    
    doc.add_paragraph('No Feedback Loop: Existing systems don\'t allow users to report incorrect detections, '
                     'preventing continuous improvement.', style='List Bullet')
    
    doc.add_paragraph(
        'The proposed project directly addresses these limitations by offering a proactive, real-time, and '
        'user-friendly solution. By integrating a hybrid detection model within a messaging platform like Telegram '
        'and providing a web-based admin dashboard, the system provides instant threat analysis directly at the '
        'point of interaction while enabling administrators to monitor system performance and user feedback. This '
        'approach significantly reduces the time-to-detection gap, empowers users with an accessible security tool, '
        'and automates a process that is currently inefficient and prone to human error.'
    )
    
    doc.add_page_break()
